import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def create_split_users_normal_table_doc():
    doc = docx.Document()

    # 1 inch margins
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)

    # Document Heading
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(6)
    run_title = p_title.add_run("3.4.3.8 Users Table Description (Role-Wise Split)")
    run_title.bold = True
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(14)

    # Intro
    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.space_after = Pt(10)
    r_desc = p_desc.add_run("The Users table is partitioned into four distinct role-based specifications, detailing the exact attributes and constraints for Citizens, Government Officials, Maintenance Contractors, and System Administrators.")
    r_desc.font.name = "Times New Roman"
    r_desc.font.size = Pt(11)

    headers = ["Column / Field Name", "Data Type", "Key Constraint", "Nullable", "Description"]
    col_widths = [1.4, 1.0, 1.3, 0.8, 2.5]

    def add_role_table(table_caption, rows_data):
        # Caption
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(10)
        p_cap.paragraph_format.space_after = Pt(4)
        r_cap = p_cap.add_run(table_caption)
        r_cap.bold = True
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(11)

        # Normal Standard Table Grid
        tbl = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
        tbl.style = 'Table Grid'
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        # Header
        hdr_cells = tbl.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
            hdr_cells[i].width = Inches(col_widths[i])
            p = hdr_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(10)
                r.font.bold = True

        # Rows
        for r_idx, row_data in enumerate(rows_data):
            row_cells = tbl.rows[r_idx + 1].cells
            for c_idx, val in enumerate(row_data):
                row_cells[c_idx].text = val
                row_cells[c_idx].width = Inches(col_widths[c_idx])
                p = row_cells[c_idx].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(9.5)
                    if c_idx == 0:
                        r.font.bold = True

        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(2)
        p_space.paragraph_format.space_after = Pt(6)

    # 1. Citizen Table
    citizen_rows = [
        ["_id", "ObjectId", "Primary Key", "No", "Unique system identifier for the citizen document in MongoDB."],
        ["name", "String", "Required", "No", "Full legal name of the registered citizen."],
        ["email", "String", "Unique Key", "No", "Email address used for citizen login authentication."],
        ["password", "String", "Required", "No", "Salted cryptographic password hash generated using bcrypt."],
        ["phone", "String", "None", "Yes", "10-digit mobile contact number for communication."],
        ["role", "String", "Default: 'Public'", "No", "Role identifier granting citizen / public access."],
        ["address", "String", "None", "Yes", "Residential address of the citizen."],
        ["createdAt", "Date", "None", "No", "Timestamp recording when the citizen account was registered."]
    ]
    add_role_table("Table 3.4.3.8.1: Citizen / Public Users Table", citizen_rows)

    # 2. Government Officials Table
    gov_rows = [
        ["_id", "ObjectId", "Primary Key", "No", "Unique system identifier for the government official."],
        ["name", "String", "Required", "No", "Full name of the municipal inspection officer."],
        ["email", "String", "Unique Key", "No", "Official government email address for login."],
        ["password", "String", "Required", "No", "Salted cryptographic password hash generated using bcrypt."],
        ["role", "String", "Value: 'Government Official'", "No", "Role identifier granting supervisory & audit access."],
        ["phone", "String", "None", "Yes", "Official mobile contact number."],
        ["department", "String", "None", "Yes", "Municipal department name (e.g., Horticulture, Civil Works)."],
        ["ward", "ObjectId", "Foreign Key (Ref: Ward)", "Yes", "Assigned municipal ward under jurisdiction."],
        ["availabilityStatus", "String", "Enum: ['Available', 'On Leave']", "No", "Real-time duty availability status for task assignment."]
    ]
    add_role_table("Table 3.4.3.8.2: Government Officials Table", gov_rows)

    # 3. Maintenance Contractors Table
    contractor_rows = [
        ["_id", "ObjectId", "Primary Key", "No", "Unique system identifier for the maintenance contractor."],
        ["name", "String", "Required", "No", "Full name of the repair contractor or technician."],
        ["email", "String", "Unique Key", "No", "Email address used for contractor portal login."],
        ["password", "String", "Required", "No", "Salted cryptographic password hash generated using bcrypt."],
        ["role", "String", "Value: 'Contractor'", "No", "Role identifier granting work order access."],
        ["phone", "String", "None", "Yes", "Contractor mobile contact number for task dispatch."],
        ["assignedParks", "Array[ObjectId]", "Foreign Key (Ref: Park)", "Yes", "List of municipal parks allocated for physical maintenance."],
        ["availabilityStatus", "String", "Enum: ['Available', 'On Leave']", "No", "Duty availability status for administrative dispatch."]
    ]
    add_role_table("Table 3.4.3.8.3: Maintenance Contractors Table", contractor_rows)

    # 4. System Administrator Table
    admin_rows = [
        ["_id", "ObjectId", "Primary Key", "No", "Unique system identifier for the administrator."],
        ["name", "String", "Required", "No", "Full name of the system administrator."],
        ["email", "String", "Unique Key", "No", "Master administrative login email address."],
        ["password", "String", "Required", "No", "Salted cryptographic master password hash (bcrypt)."],
        ["role", "String", "Value: 'Admin'", "No", "Grants system-wide administrative control and master privileges."],
        ["createdAt", "Date", "None", "No", "Timestamp recording when the admin account was created."]
    ]
    add_role_table("Table 3.4.3.8.4: System Administrator Table", admin_rows)

    output_path = r'd:\park_monitoring_system08\park-management\Users_Split_Normal_Tables.docx'
    doc.save(output_path)
    print("Split Normal Table DOCX successfully created at:", output_path)

if __name__ == '__main__':
    create_split_users_normal_table_doc()
