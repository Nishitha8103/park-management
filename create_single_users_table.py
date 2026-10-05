import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def create_single_users_table_doc():
    doc = docx.Document()

    # 1 inch margins
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)

    # Heading
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("3.4.3.8 Users Table Description")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)

    # Paragraph Description
    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.space_after = Pt(10)
    r_desc = p_desc.add_run("The Users Table stores the core authentication credentials, user identity profiles, contact details, and role classifications across Citizens, Government Officials, Maintenance Contractors, and System Administrators.")
    r_desc.font.name = "Times New Roman"
    r_desc.font.size = Pt(11)

    # Table Title
    p_tbl = doc.add_paragraph()
    p_tbl.paragraph_format.space_before = Pt(6)
    p_tbl.paragraph_format.space_after = Pt(4)
    r_tbl = p_tbl.add_run("Table 3.4.3.8: Structure of Users Table")
    r_tbl.bold = True
    r_tbl.font.name = "Times New Roman"
    r_tbl.font.size = Pt(11)

    headers = ["Column / Field Name", "Data Type", "Key Constraint", "Nullable", "Default Value", "Description"]
    col_widths = [1.3, 0.9, 1.1, 0.7, 0.9, 2.1]

    rows_data = [
        ["_id", "ObjectId", "Primary Key", "No", "Auto-generated", "Unique system identifier for the user document in MongoDB."],
        ["name", "String", "Required", "No", "None", "Full legal name of the registered user."],
        ["email", "String", "Unique Key", "No", "None", "Unique email address utilized for login authentication."],
        ["password", "String", "Required", "No", "None", "Salted cryptographic password hash generated using bcrypt."],
        ["role", "String", "Check / Enum", "No", "'Public'", "User access role ('Admin', 'Contractor', 'Government Official', 'Public')."],
        ["phone", "String", "None", "Yes", "null", "10-digit primary mobile contact number for communication and alerts."],
        ["address", "String", "None", "Yes", "''", "Residential or business physical address of the user."],
        ["department", "String", "None", "Yes", "null", "Municipal department name (for Government Officials)."],
        ["ward", "ObjectId", "Foreign Key", "Yes", "null", "Assigned municipal ward reference under jurisdiction."],
        ["assignedParks", "Array[ObjectId]", "Foreign Key", "Yes", "[]", "Public parks allocated to official or contractor."],
        ["availabilityStatus", "String", "Check / Enum", "No", "'Available'", "Real-time duty availability status ('Available', 'On Leave', 'Unavailable')."],
        ["createdAt", "Date", "None", "No", "Date.now", "Automatic timestamp recording when the user account was registered."],
        ["updatedAt", "Date", "None", "No", "Date.now", "Automatic timestamp recording the latest profile or credential update."]
    ]

    # Create Normal Standard Word Table (Table Grid style)
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].width = Inches(col_widths[i])
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            r.font.bold = True

    # Data Rows
    for r_idx, row_data in enumerate(rows_data):
        row_cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = val
            row_cells[c_idx].width = Inches(col_widths[c_idx])
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(9.5)
                if c_idx == 0:
                    r.font.bold = True

    output_path = r'd:\park_monitoring_system08\park-management\Users_Table_Description_Normal.docx'
    doc.save(output_path)
    print("Normal table DOCX successfully saved at:", output_path)

if __name__ == '__main__':
    create_single_users_table_doc()
