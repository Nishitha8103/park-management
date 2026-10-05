import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

OUTPUT_DIR = r"d:\park_monitoring_system08\park-management\documentation_chapters_v2"
COMBINED_FILE = os.path.join(OUTPUT_DIR, "Complete_Park_Monitoring_Project_Report.docx")

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_full_report():
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1)
        
        header = section.header
        hp = header.paragraphs[0]
        hp.text = "Park Monitoring & Management System"
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hp.style.font.name = 'Times New Roman'
        hp.style.font.size = Pt(9)
        hp.style.font.color.rgb = RGBColor(120, 120, 120)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = "Department of MCA, PIM, Udupi"
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp.style.font.name = 'Times New Roman'
        fp.style.font.size = Pt(9)
        fp.style.font.color.rgb = RGBColor(120, 120, 120)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(30, 30, 30)
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_after = Pt(6)

    def add_h1(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(12)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)

    def add_h2(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)

    def add_h3(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)

    def add_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(6)
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.bold = True
            rb.font.name = 'Times New Roman'
            rb.font.size = Pt(12)
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)

    # ------------------ PRELIMINARY ------------------
    add_h1("MANGALORE UNIVERSITY")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Project Report on\n“PARK MONITORING SYSTEM”\n\nCarried out and submitted by\n")
    r.font.bold = True
    r.font.size = Pt(14)
    r2 = p.add_run("NISHITHA\nReg. No.: P05PP23S126010\nIV Semester MCA Student\nPoornaprajna Institute of Management, Udupi\n\n")
    r2.font.bold = True
    r2.font.size = Pt(13)
    r3 = p.add_run("Under the guidance of\nProf. Venugopala Rao A. S.\nHOD\nDept. of MCA, PIM, Udupi.\n\nIn the partial fulfilment of the requirement for the award of degree in Master of Computer Applications during the academic year 2024-25\n\nPoornaprajna Institute of Management,\nUdupi\n2024-2025\n")
    r3.font.size = Pt(12)

    doc.add_page_break()
    add_h1("CERTIFICATE")
    add_p("This is to certify that the project entitled “Park Monitoring System” has been carried out by Nishitha (Reg.No.: P05PP23S126010), student of fourth semester MCA (Masters of Computer Applications) under the supervision of Prof. Venugopala Rao A.S., MCA Department, Poornaprajna Institute of Management, Udupi. The project is submitted in partial fulfilment of the requirement for the award of Master of Computer Applications by Mangalore University during the Academic year 2024-2025.\n\n\n\n")
    p_sig = doc.add_paragraph()
    p_sig.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_sig.add_run("Internal Guide\t\t\t\t\tHead of the Department\n\n\n\nInternal Examiner\t\t\t\t\tExternal Examiner\n\nSubmitted for the viva-voce examination held on: ..................")

    doc.add_page_break()
    add_h1("DECLARATION")
    add_p("This project work entitled “PARK MONITORING SYSTEM” has been successfully carried out by me under the supervision and guidance of Prof. Venugopala Rao A.S., Head of the MCA Department, Poornaprajna Institute of Management, Udupi. This project is submitted in partial fulfilment for the award of Masters of Computer Application degree by Mangalore University during the academic year 2024–25. This work or any part of this work has not been submitted to any other university or Institute/School for the award of any other Degree or Diploma.\n\n\nDate: ...................\nPlace: Udupi\t\t\t\t\tName: NISHITHA\n\t\t\t\t\t\tReg. No.: P05PP23S126010")

    doc.add_page_break()
    add_h1("ACKNOWLEDGEMENT")
    add_p("I take this opportunity to express my sincere thanks to my guide Prof. Venugopala Rao A. S., Head of the MCA Department, Poornaprajna Institute of Management, for his all-round guidance, timely help at every stage of this project and for the valuable suggestions and unlimited support.\n\nI would like to express my gratitude to Dr. P. S. Aithal, Director, Poornaprajna Institute of Management, for granting the necessary permissions and providing all kinds of infrastructure facilities in the department to carry out this project successfully.\n\nAll faculty members and non-teaching staff of the MCA department.\n\nI would like to express my heartfelt gratitude to my parents, friends and well-wishers who have always inspired and blessed me, including those whom I may have inadvertently failed to mention. Above all, with all my heart, I thank you God for empowering me with the dedication, focus and patience to carry out this project successfully.\n\n\nNISHITHA\nReg. No.: P05PP23S126010")

    doc.add_page_break()
    add_h1("ABSTRACT")
    add_p("Urban green space management and public park maintenance are critical challenges faced by rapidly growing cities like Bengaluru. The increasing wear and tear of park amenities, damaged play equipment, broken benches, non-functional lighting, overflowing waste receptacles, and neglected garden flora have become a major concern for the Bruhat Bengaluru Mahanagara Palike (BBMP). These neglected park issues not only affect the city's cleanliness and natural beauty but also pose severe safety and environmental risks to visiting citizens and children. To address these issues, there is a pressing need for a comprehensive digital platform that enables timely reporting, monitoring, and resolution of park-related problems through active citizen participation and structured administrative coordination. This project presents the development of a Progressive Web Application (PWA) aimed at supporting BBMP's urban green management initiative. The system provides a real-time, transparent, and user-friendly platform for reporting, tracking, and resolving park maintenance complaints across Bengaluru. It bridges the communication gap between Citizens, Field Officials, Maintenance Contractors, and Administrators by ensuring smooth information flow and coordinated action.")
    add_p("The system comprises four major integrated modules: Citizens, Admin, Government Officials, and Maintenance Contractors. The Citizens module enables citizens to register, log in, update profiles, explore park directories with amenity details, book commercial stall slots with digital payments, register for community events, and raise maintenance complaints by capturing live photos through their device camera along with automatic GPS location coordinates and problem descriptions. Citizens can also check the live status of their complaints and track the complete resolution lifecycle. Additionally, the system sends automated email notifications to citizens upon successful complaint registration, assignment, and final resolution, ensuring complete transparency throughout the process.")
    add_p("The Admin module acts as the central control panel, managing all administrative activities such as complaint assignments, contractor oversight, official supervision, stall booking reviews, event registrations, and report generation. It includes sub-modules for managing users, officials, contractors, districts, corporations, zones, and wards, allowing the administrator to maintain proper geographic and operational control across the entire city. When a complaint is assigned or resolved, automated email alerts are triggered to both the assigned field staff and the concerned citizen, improving communication efficiency and accountability.")
    add_p("The Government Officials and Contractors modules empower on-ground field staff to log in, view assigned tasks, submit daily work progress logs with photographic evidence, request necessary maintenance materials, verify completed repairs on-site, and maintain a complete work history for performance auditing. By integrating these modules with real-time email notification support, interactive map visualization, and PWA capabilities, the application ensures seamless communication, efficient complaint resolution, and improved park management workflows.")
    add_p("Keywords: Park Monitoring, PWA, Proactive, Transparency, Complaints, Maintenance, BBMP", "")

    # Table of Contents
    doc.add_page_break()
    add_h1("TABLE OF CONTENTS")
    toc_data = [
        ("1.", "Introduction", "1-11"),
        ("", "1.1 Introduction", "1"),
        ("", "1.2 Overview of the project", "2"),
        ("", "1.3 Problem Statement", "3"),
        ("", "1.4 Motivation", "4"),
        ("", "1.5 Significance of the Study", "5"),
        ("", "1.6 Objectives", "7"),
        ("", "1.7 Scope of the Project", "7"),
        ("", "1.8 Features", "10"),
        ("2.", "Literature Review", "12-16"),
        ("", "2.1 Introduction", "12"),
        ("", "2.2 Result analysis", "12"),
        ("", "2.3 Identified gaps in the Literature", "15"),
        ("", "2.4 Existing System", "16"),
        ("", "2.5 Proposed System", "16"),
        ("3.", "System Analysis", "17-23"),
        ("", "3.1 Introduction", "17"),
        ("", "3.2 Overall Descriptions", "18"),
        ("", "3.3 Specific Requirements", "19"),
        ("", "3.4 Functional Requirements", "20"),
        ("", "3.5 Performance Requirements", "22"),
        ("", "3.6 Design Constraints", "23"),
        ("", "3.7 Other Requirements", "23"),
        ("", "3.8 Safety Requirements", "23"),
        ("", "3.9 Security Requirements", "23"),
        ("4.", "Design And Methodology", "24-36"),
        ("", "4.1 System Design", "24"),
        ("", "4.2 Detailed Design", "28"),
        ("", "4.3 Database Design", "34"),
        ("5.", "Implementation Details", "37-46"),
        ("", "5.1 Introduction", "37"),
        ("", "5.2 Hardware And Software tools used", "37"),
        ("", "5.3 Source Code", "38"),
        ("6.", "Result and Evaluation", "47-51"),
        ("", "6.1 Introduction", "47"),
        ("", "6.2 Test Scenario", "47"),
        ("", "6.3 Test Cases", "47"),
        ("7.", "Conclusion and Future Work", "52-53"),
        ("", "7.1 Conclusion", "52"),
        ("", "7.2 Future Work", "52"),
        ("", "Appendices", "54-64"),
        ("", "References", "65")
    ]
    t_table = doc.add_table(rows=1, cols=3)
    t_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_table.rows[0].cells[0].text = "Chapter No."
    t_table.rows[0].cells[1].text = "Chapter Name"
    t_table.rows[0].cells[2].text = "Page No."
    for cell in t_table.rows[0].cells:
        cell.paragraphs[0].runs[0].font.bold = True
        set_cell_background(cell, "E2E8F0")

    for row in toc_data:
        r_cells = t_table.add_row().cells
        r_cells[0].text = row[0]
        r_cells[1].text = row[1]
        r_cells[2].text = row[2]

    doc.save(COMBINED_FILE)
    print("Generated Consolidated Full Report:", COMBINED_FILE)

if __name__ == "__main__":
    create_full_report()
