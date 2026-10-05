import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_users_table_doc():
    doc = docx.Document()

    # Page setup - 1 inch margins
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)

    # Style helpers
    def set_cell_background(cell, fill_color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def add_custom_table(doc, title, headers, rows, col_widths):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(title)
        run.bold = True
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0, 51, 102)

        table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        # Header Row
        hdr_cells = table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
            hdr_cells[i].width = Inches(col_widths[i])
            set_cell_background(hdr_cells[i], '1E3A8A')
            set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
            p = hdr_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)

        # Data Rows
        for r_idx, row_data in enumerate(rows):
            row_cells = table.rows[r_idx + 1].cells
            bg_color = 'F8FAFC' if r_idx % 2 == 1 else 'FFFFFF'
            for c_idx, val in enumerate(row_data):
                row_cells[c_idx].text = val
                row_cells[c_idx].width = Inches(col_widths[c_idx])
                set_cell_background(row_cells[c_idx], bg_color)
                set_cell_margins(row_cells[c_idx], top=90, bottom=90, left=140, right=140)
                p = row_cells[c_idx].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(10)
                    if c_idx == 0:
                        r.font.bold = True

        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(4)
        p_space.paragraph_format.space_after = Pt(8)

    # Document Heading
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(12)
    r = title_p.add_run('3.4.3.8 Users Table Description (Role-Separated)')
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(0, 51, 102)

    intro_p = doc.add_paragraph()
    intro_p.paragraph_format.space_after = Pt(10)
    r_intro = intro_p.add_run('The Users Table is divided into four distinct role-based specifications, detailing the primary attributes, data types, constraints, and descriptions governing Citizens, Government Officials, Maintenance Contractors, and System Administrators.')
    r_intro.font.name = 'Times New Roman'
    r_intro.font.size = Pt(11)

    headers = ['Field Name', 'Data Type', 'Key Constraint', 'Nullable', 'Description']
    col_widths = [1.2, 0.9, 1.2, 0.8, 2.4]

    # 1. Citizen Table
    rows_citizen = [
        ['_id', 'ObjectId', 'Primary Key', 'No', 'Unique identifier for the citizen document in MongoDB.'],
        ['name', 'String', 'Required', 'No', 'Full legal name of the registered citizen.'],
        ['email', 'String', 'Unique Key', 'No', 'Email address used for citizen login authentication.'],
        ['password', 'String', 'Required', 'No', 'Salted cryptographic password hash generated via bcrypt.'],
        ['phone', 'String', 'None', 'Yes', '10-digit mobile contact number for notifications.'],
        ['role', 'String', 'Default: "Public"', 'No', 'Access control role identifier for public user services.'],
        ['createdAt', 'Date', 'None', 'No', 'Timestamp recording when the citizen account was registered.']
    ]
    add_custom_table(doc, '1. Citizen / Public Users Table', headers, rows_citizen, col_widths)

    # 2. Government Officials Table
    rows_official = [
        ['_id', 'ObjectId', 'Primary Key', 'No', 'Unique identifier for the government official.'],
        ['name', 'String', 'Required', 'No', 'Full name of the municipal inspection official.'],
        ['email', 'String', 'Unique Key', 'No', 'Official municipal email address used for login.'],
        ['password', 'String', 'Required', 'No', 'Salted cryptographic password hash generated via bcrypt.'],
        ['role', 'String', 'Value: "Government Official"', 'No', 'Role identifier granting supervisory & audit access.'],
        ['phone', 'String', 'None', 'Yes', 'Official contact phone number.'],
        ['department', 'String', 'None', 'Yes', 'Municipal department name (e.g., Horticulture, Civil Works).'],
        ['ward', 'ObjectId', 'Foreign Key (Ref: Ward)', 'Yes', 'Assigned municipal ward under official jurisdiction.'],
        ['availabilityStatus', 'String', 'Enum: ["Available", "On Leave"]', 'No', 'Real-time duty availability status for task assignment.']
    ]
    add_custom_table(doc, '2. Government Officials Table', headers, rows_official, col_widths)

    # 3. Maintenance Contractors Table
    rows_contractor = [
        ['_id', 'ObjectId', 'Primary Key', 'No', 'Unique identifier for the maintenance contractor.'],
        ['name', 'String', 'Required', 'No', 'Full name of the repair contractor or technician.'],
        ['email', 'String', 'Unique Key', 'No', 'Email address utilized for contractor portal login.'],
        ['password', 'String', 'Required', 'No', 'Salted cryptographic password hash generated via bcrypt.'],
        ['role', 'String', 'Value: "Contractor"', 'No', 'Role identifier granting work order access.'],
        ['phone', 'String', 'None', 'Yes', 'Contractor mobile contact number for work dispatch.'],
        ['assignedParks', 'Array[ObjectId]', 'Foreign Key (Ref: Park)', 'Yes', 'List of municipal parks allocated for physical repairs.'],
        ['availabilityStatus', 'String', 'Enum: ["Available", "On Leave"]', 'No', 'Work availability status for administrative dispatch.']
    ]
    add_custom_table(doc, '3. Maintenance Contractors Table', headers, rows_contractor, col_widths)

    # 4. Admin Table
    rows_admin = [
        ['_id', 'ObjectId', 'Primary Key', 'No', 'Unique system identifier for the administrator.'],
        ['name', 'String', 'Required', 'No', 'Full name of the system administrator.'],
        ['email', 'String', 'Unique Key', 'No', 'Master administrative login email address.'],
        ['password', 'String', 'Required', 'No', 'Salted cryptographic master password hash (bcrypt).'],
        ['role', 'String', 'Value: "Admin"', 'No', 'Grants system-wide administrative control and master privileges.'],
        ['createdAt', 'Date', 'None', 'No', 'Timestamp recording when the admin account was created.']
    ]
    add_custom_table(doc, '4. System Administrator Table', headers, rows_admin, col_widths)

    output_path = r'd:\park_monitoring_system08\park-management\Users_Table_Description.docx'
    doc.save(output_path)
    print('DOCX successfully created at:', output_path)

if __name__ == '__main__':
    create_users_table_doc()
