import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def generate_database_tables_docx():
    doc = docx.Document()

    # 1 inch margins
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(6)
    r_title = p_title.add_run("3.4.3 Database Table Descriptions (Data Dictionary)")
    r_title.bold = True
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(14)

    # Document Intro
    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.space_after = Pt(10)
    r_desc = p_desc.add_run("This section presents the database table structures and data dictionaries for all key entities implemented in the Park Monitoring System database, listing the essential attributes, data types, constraints, and operational descriptions.")
    r_desc.font.name = "Times New Roman"
    r_desc.font.size = Pt(11)

    headers = ["Column / Field Name", "Data Type", "Key Constraint", "Nullable", "Description"]
    col_widths = [1.4, 1.0, 1.3, 0.8, 2.5]

    def add_table_section(table_caption, rows_data):
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(10)
        p_cap.paragraph_format.space_after = Pt(4)
        r_cap = p_cap.add_run(table_caption)
        r_cap.bold = True
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(11)

        tbl = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
        tbl.style = 'Table Grid'
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False

        # Header
        hdr_cells = tbl.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
            hdr_cells[i].width = Inches(col_widths[i])
            p = hdr_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(10)
                r.font.bold = True

        # Rows
        for r_idx, row_data in enumerate(rows_data):
            row_cells = tbl.rows[r_idx + 1].cells
            for c_idx, val in enumerate(row_data):
                row_cells[c_idx].text = val
                row_cells[c_idx].width = Inches(col_widths[c_idx])
                p = row_cells[c_idx].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(9.5)
                    if c_idx == 0:
                        r.font.bold = True

        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(2)
        p_space.paragraph_format.space_after = Pt(6)

    # 1. Parks Table
    park_rows = [
        ["_id", "ObjectId", "Primary Key", "No", "Unique system identifier for the park document."],
        ["name", "String", "Required", "No", "Official name of the municipal public park."],
        ["parkCode", "String", "Required, Unique", "No", "Unique alphanumeric administrative code for the park."],
        ["zone", "ObjectId", "Foreign Key (Ref: Zone)", "No", "Municipal zone reference under Bangalore jurisdiction."],
        ["ward", "ObjectId", "Foreign Key (Ref: Ward)", "No", "Municipal ward reference."],
        ["latitude", "String", "None", "Yes", "Geographical latitude coordinate for mapping."],
        ["longitude", "String", "None", "Yes", "Geographical longitude coordinate for mapping."],
        ["facilities", "Array[String]", "None", "Yes", "List of amenities (Tracks, Play Area, Gym, Lights, Benches, Restrooms)."],
        ["contractor", "ObjectId", "Foreign Key (Ref: User)", "Yes", "Assigned maintenance contractor for the park."],
        ["governmentOfficial", "ObjectId", "Foreign Key (Ref: User)", "Yes", "Assigned supervising municipal inspection official."],
        ["status", "String", "Enum: ['Active', 'Closed']", "No", "Operational status of the park (Default: 'Active')."]
    ]
    add_table_section("Table 3.4.3.1: Structure of Parks Table", park_rows)

    # 2. Complaints Table
    complaint_rows = [
        ["_id", "ObjectId", "Primary Key", "No", "Unique identifier for the complaint record."],
        ["complaintNumber", "String", "Required, Unique", "No", "Unique tracking ticket ID (e.g., TKT-2026-1042)."],
        ["user", "ObjectId", "Foreign Key (Ref: User)", "No", "Citizen who lodged the maintenance complaint."],
        ["park", "ObjectId", "Foreign Key (Ref: Park)", "No", "Park where the infrastructure issue is located."],
        ["category", "String", "Required", "No", "Defect type (Equipment, Lighting, Sanitation, Benches, Trees)."],
        ["description", "String", "Required", "No", "Textual description of the reported problem."],
        ["images", "Array[String]", "Required", "No", "Live camera captured photo URLs with embedded GPS watermarks."],
        ["status", "String", "Check / Enum", "No", "Lifecycle state (New, Assigned, In Progress, Inspection Pending, Closed)."],
        ["assignedContractor", "ObjectId", "Foreign Key (Ref: Contractor)", "Yes", "Contractor assigned to execute physical repairs."],
        ["assignedOfficial", "ObjectId", "Foreign Key (Ref: User)", "Yes", "Government official assigned to audit and verify repair."],
        ["afterImages", "Array[String]", "None", "Yes", "Post-repair photo proof uploaded by the contractor."],
        ["inspectionRemarks", "String", "None", "Yes", "Official audit notes and digital sign-off comments."],
        ["createdAt", "Date", "None", "No", "Timestamp recording when the grievance was lodged."]
    ]
    add_table_section("Table 3.4.3.2: Structure of Complaints Table", complaint_rows)

    # 3. Stall Bookings Table
    stall_rows = [
        ["_id", "ObjectId", "Primary Key", "No", "Unique commercial stall booking transaction identifier."],
        ["park", "ObjectId", "Foreign Key (Ref: Park)", "No", "Park where the commercial stall slot is located."],
        ["user", "ObjectId", "Foreign Key (Ref: User)", "No", "Citizen / Vendor who booked the stall slot."],
        ["stallName", "String", "Required", "No", "Commercial trade or stall name."],
        ["productsType", "String", "Required", "No", "Category of products or services offered at the stall."],
        ["amountPaid", "Number", "Required", "No", "Total booking amount paid digitally."],
        ["status", "String", "Enum: ['Pending', 'Confirmed', 'Rejected']", "No", "Booking confirmation status."],
        ["bookingDate", "Date", "None", "No", "Date and time when the stall slot is booked."]
    ]
    add_table_section("Table 3.4.3.3: Structure of Stall Bookings Table", stall_rows)

    # 4. Events Table
    event_rows = [
        ["_id", "ObjectId", "Primary Key", "No", "Unique identifier for the community event."],
        ["title", "String", "Required", "No", "Title of the eco-friendly event (e.g., Tree Plantation, Yoga)."],
        ["description", "String", "Required", "No", "Detailed description and guidelines of the event."],
        ["parkId", "ObjectId", "Foreign Key (Ref: Park)", "Yes", "Venue park where the event is scheduled."],
        ["eventDate", "Date", "Required", "No", "Date and scheduled time of the event."],
        ["capacity", "Number", "Required", "No", "Maximum participant attendance capacity."],
        ["isActive", "Boolean", "None", "No", "Active status flag indicating if registration is open."]
    ]
    add_table_section("Table 3.4.3.4: Structure of Events Table", event_rows)

    # 5. Material Requests Table
    mat_rows = [
        ["_id", "ObjectId", "Primary Key", "No", "Unique material request identifier."],
        ["contractor", "ObjectId", "Foreign Key (Ref: Contractor)", "No", "Contractor requesting repair tools/parts."],
        ["park", "ObjectId", "Foreign Key (Ref: Park)", "Yes", "Park where materials are required."],
        ["materialName", "String", "Required", "No", "Name of the requested equipment or spare part."],
        ["quantity", "Number", "Required", "No", "Quantity of materials requested."],
        ["reason", "String", "Required", "No", "Operational reason for material requisition."],
        ["status", "String", "Enum: ['Pending', 'Approved', 'Rejected']", "No", "Administrative approval status (Default: 'Pending')."]
    ]
    add_table_section("Table 3.4.3.5: Structure of Material Requests Table", mat_rows)

    # 6. Leave Requests Table
    leave_rows = [
        ["_id", "ObjectId", "Primary Key", "No", "Unique leave application identifier."],
        ["leaveId", "String", "Required, Unique", "No", "Unique tracking reference ID for the leave request."],
        ["applicantId", "String", "Required", "No", "User ID of the applying contractor or government official."],
        ["applicantRole", "String", "Required", "No", "Role of applicant ('contractor', 'government_official')."],
        ["leaveType", "String", "Required", "No", "Type of leave (Casual Leave, Medical Leave, Emergency)."],
        ["startDate", "Date", "Required", "No", "Starting date of the leave period."],
        ["endDate", "Date", "Required", "No", "Ending date of the leave period."],
        ["reason", "String", "Required", "No", "Reason for leave request."],
        ["status", "String", "Enum: ['Pending', 'Approved', 'Rejected']", "No", "Administrative leave decision status."]
    ]
    add_table_section("Table 3.4.3.6: Structure of Leave Requests Table", leave_rows)

    # 7. Feedback Table
    feedback_rows = [
        ["_id", "ObjectId", "Primary Key", "No", "Unique feedback submission identifier."],
        ["feedbackId", "String", "Required, Unique", "No", "Unique reference ID for the feedback record."],
        ["user", "ObjectId", "Foreign Key (Ref: User)", "Yes", "Citizen who submitted the review."],
        ["park", "ObjectId", "Foreign Key (Ref: Park)", "No", "Park for which the feedback is submitted."],
        ["overallRating", "Number", "Required (1 to 5)", "No", "Overall star rating given by the citizen."],
        ["cleanlinessRating", "Number", "Default: 4", "No", "Cleanliness rating score (1 to 5 scale)."],
        ["maintenanceRating", "Number", "Default: 4", "No", "Infrastructure maintenance score (1 to 5 scale)."],
        ["comments", "String", "None", "Yes", "Qualitative review and remarks submitted by the visitor."]
    ]
    add_table_section("Table 3.4.3.7: Structure of Feedback Table", feedback_rows)

    output_path = r'd:\park_monitoring_system08\park-management\Database_Tables_Description.docx'
    doc.save(output_path)
    print("Database Tables DOCX successfully created at:", output_path)

if __name__ == '__main__':
    generate_database_tables_docx()
