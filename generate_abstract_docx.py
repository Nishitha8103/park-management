import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_abstract_docx(filename):
    doc = Document()
    
    # 1 inch margins
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # Heading: ABSTRACT
    h_p = doc.add_paragraph()
    h_p.paragraph_format.space_before = Pt(12)
    h_p.paragraph_format.space_after = Pt(14)
    h_p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    r_h = h_p.add_run("ABSTRACT")
    r_h.font.name = 'Times New Roman'
    r_h.font.size = Pt(14)
    r_h.font.bold = True
    r_h.font.color.rgb = RGBColor(0, 0, 0)
    
    abstract_text = (
        "Urban green spaces and public municipal parks play an indispensable role in maintaining environmental sustainability, "
        "promoting public wellness, and improving the quality of urban life. However, modern urban park management faces critical "
        "operational hurdles, including delayed grievance resolution, lack of geospatial verification in citizen reports, untracked "
        "contractor field work, and the absence of institutional multi-tier verification before public work sign-offs. Conventional "
        "complaint systems frequently suffer from fraudulent or inaccurate image submissions and manual coordination silos among municipal "
        "administrators, field contractors, and government monitoring officials. To overcome these limitations, this project presents the "
        "Park Monitoring System—a comprehensive, full-stack digital governance platform engineered with the MERN (MongoDB, Express.js, React.js, "
        "Node.js) technology stack.\n\n"
        "The proposed platform introduces an authentic crowdsourced incident reporting pipeline for citizens, enforcing real-time in-app camera "
        "capture paired with GPS geocoding and timestamp watermarking while strictly disabling gallery uploads. For administrative governance, the "
        "system deploys a dual-assignment workflow where incoming complaints are simultaneously assigned to specialized field contractors and "
        "designated government inspection officials under strict Service-Level Agreement (SLA) deadlines. Field contractors manage digital work "
        "orders, request necessary maintenance materials, and upload verifiable photographic proof of resolution. Crucially, before any issue is "
        "marked resolved, government officials perform an on-site physical audit within the park's geofenced boundaries and execute a digital "
        "sign-off. In addition, the system integrates facility reservations (commercial stalls and event venues), automated staff leave management, "
        "and multi-dimensional analytical audit reporting. The resulting system ensures tamper-proof accountability, reduces administrative overhead, "
        "eliminates bogus grievance reporting, and establishes a transparent, data-driven operational framework for smart municipal governance."
    )
    
    keywords_text = "Keywords: Smart City Governance, Park Monitoring System, MERN Stack, GPS Watermarking, Dual-Assignment Workflow, Digital Sign-off, Facility Reservation."

    for paragraph_str in abstract_text.split("\n\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        run = p.add_run(paragraph_str)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0, 0, 0)
        
    p_kw = doc.add_paragraph()
    p_kw.paragraph_format.space_before = Pt(14)
    p_kw.paragraph_format.space_after = Pt(6)
    p_kw.paragraph_format.line_spacing = 1.5
    p_kw.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    r_kw_title = p_kw.add_run("Keywords: ")
    r_kw_title.font.name = 'Times New Roman'
    r_kw_title.font.size = Pt(11)
    r_kw_title.font.bold = True
    r_kw_title.font.color.rgb = RGBColor(0, 0, 0)
    
    r_kw = p_kw.add_run("Smart City Governance, Park Monitoring System, MERN Stack, GPS Watermarking, Dual-Assignment Workflow, On-Site Physical Inspection, Digital Sign-Off, Urban Civic Maintenance.")
    r_kw.font.name = 'Times New Roman'
    r_kw.font.size = Pt(11)
    r_kw.font.italic = True
    r_kw.font.color.rgb = RGBColor(0, 0, 0)

    doc.save(filename)
    print(f"Abstract docx created at: {filename}")

if __name__ == "__main__":
    create_abstract_docx(r"d:\park_monitoring_system08\park-management\Project_Abstract.docx")
