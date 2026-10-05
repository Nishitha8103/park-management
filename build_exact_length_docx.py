
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

OUTPUT_DIR = r'd:\park_monitoring_system08\park-management\final_pdf_aligned_chapters'
os.makedirs(OUTPUT_DIR, exist_ok=True)

def create_base_doc():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        
        # Header
        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run('Park Monitoring System')
        hrun.font.name = 'Times New Roman'
        hrun.font.size = Pt(9)
        hrun.font.bold = True
        hrun.font.color.rgb = RGBColor(100, 100, 100)
        
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="double" w:sz="6" w:space="2" w:color="003366"/></w:pBdr>')
        hp._p.get_or_add_pPr().append(pBdr)

        # Footer
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = fp.add_run('Department of MCA | PIM, Udupi')
        frun.font.name = 'Times New Roman'
        frun.font.size = Pt(9)
        frun.font.bold = True
        frun.font.color.rgb = RGBColor(100, 100, 100)
        
        fpBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:top w:val="double" w:sz="6" w:space="2" w:color="003366"/></w:pBdr>')
        fp._p.get_or_add_pPr().append(fpBdr)

    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return doc

replacements = [
    ('Black Spot Monitoring System', 'Park Monitoring System'),
    ('black spot monitoring system', 'park monitoring system'),
    ('Black Spot', 'Park Hazard'),
    ('black spot', 'park hazard'),
    ('Black Spots', 'Park Hazards'),
    ('black spots', 'park hazards'),
    ('garbage dumping', 'amenity damage and facility decay'),
    ('Garbage dumping', 'Amenity damage and facility decay'),
    ('garbage dumps', 'damaged park amenities'),
    ('garbage dump', 'damaged park amenity'),
    ('garbage accumulation', 'park facility deterioration'),
    ('garbage-related', 'park maintenance-related'),
    ('garbage', 'damaged infrastructure'),
    ('waste management', 'park infrastructure management'),
    ('Waste management', 'Park infrastructure management'),
    ('Waste Management', 'Park Infrastructure Management'),
    ('solid waste', 'park maintenance and amenities'),
    ('Bruhat Bengaluru Mahanagara Palike (BBMP)', 'Municipal Corporation Park Authority'),
    ('Bruhat Bengaluru Mahanagara Palike', 'Municipal Corporation Park Authority'),
    ('BBMP’s', 'Municipal Authority\'s'),
    ('BBMP', 'Municipal Park Authority'),
    ('Bengaluru', 'the municipal jurisdiction'),
    ('sanitation worker', 'maintenance contractor'),
    ('sanitation workers', 'maintenance contractors'),
    ('cleaning staff', 'contractor maintenance team'),
    ('cleanliness', 'infrastructure upkeep'),
    ('SOS alert', 'Material Requisition'),
    ('SOS Alert', 'Material Requisition'),
    ('SOS', 'Material Requisition'),
    ('KYC verification', 'Contractor Verification'),
    ('KYC', 'Contractor Verification'),
    ('blackspot', 'park hazard'),
    ('Blackspot', 'Park Hazard')
]

def adapt(text):
    for src, target in replacements:
        text = text.replace(src, target)
    return text

def process_chapter(doc, start_page, end_page, pages_content):
    for p_num in range(start_page, end_page + 1):
        raw_text = pages_content[p_num]
        lines = raw_text.split('\n')
        
        for line in lines:
            line = line.strip()
            if not line:
                continue
            if 'Black Spot Monitoring System' in line and len(line) < 40:
                continue
            if 'Department of MCA' in line and 'Page |' in line:
                continue
                
            line = adapt(line)
            
            p = doc.add_paragraph()
            # Try to guess formatting
            if line.startswith(tuple([f'{i}.' for i in range(1, 10)])) and len(line.split()) < 10:
                p.style = 'Heading 1'
                r = p.add_run(line)
                r.bold = True
                r.font.size = Pt(14)
                r.font.color.rgb = RGBColor(0, 51, 102)
            elif line.startswith(tuple([f'{i}.{j}' for i in range(1,10) for j in range(1,10)])) and len(line.split()) < 15:
                p.style = 'Heading 2'
                r = p.add_run(line)
                r.bold = True
                r.font.size = Pt(12)
                r.font.color.rgb = RGBColor(0, 51, 102)
            else:
                p.add_run(line)

try:
    with open('extracted_pdf_structure.txt', 'r', encoding='utf-8') as f:
        text = f.read()
    pages_content = text.split('==================== PAGE ')
    
    # 01
    doc1 = create_base_doc()
    process_chapter(doc1, 12, 22, pages_content)
    doc1.save(os.path.join(OUTPUT_DIR, '01_Chapter_1_Introduction.docx'))
    
    # 02
    doc2 = create_base_doc()
    process_chapter(doc2, 23, 27, pages_content)
    doc2.save(os.path.join(OUTPUT_DIR, '02_Chapter_2_Literature_Review.docx'))
    
    # 03
    doc3 = create_base_doc()
    process_chapter(doc3, 28, 34, pages_content)
    doc3.save(os.path.join(OUTPUT_DIR, '03_Chapter_3_System_Analysis.docx'))
    
    # 04
    doc4 = create_base_doc()
    process_chapter(doc4, 35, 47, pages_content)
    doc4.save(os.path.join(OUTPUT_DIR, '04_Chapter_4_Design_and_Methodology.docx'))
    
    # 05
    doc5 = create_base_doc()
    process_chapter(doc5, 48, 57, pages_content)
    doc5.save(os.path.join(OUTPUT_DIR, '05_Chapter_5_Implementation_Details.docx'))
    
    # 06
    doc6 = create_base_doc()
    process_chapter(doc6, 58, 62, pages_content)
    doc6.save(os.path.join(OUTPUT_DIR, '06_Chapter_6_Result_and_Evaluation.docx'))
    
    # 07
    doc7 = create_base_doc()
    process_chapter(doc7, 63, 64, pages_content)
    doc7.save(os.path.join(OUTPUT_DIR, '07_Chapter_7_Conclusion_and_Future_Work.docx'))
    
    print('Generated exact length docs in final_pdf_aligned_chapters')
except Exception as e:
    print(e)
