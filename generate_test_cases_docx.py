import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders_xml = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
        f'  <w:left w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
        f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
        f'  <w:insideH w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
        f'  <w:insideV w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders_xml)

def create_test_cases_docx(filename):
    doc = Document()
    
    # Page setup - 1 inch margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        
    modules = [
        {
            "num": "5.6.1",
            "title": "User Registration Form",
            "cases": [
                ("1.", "If user clicks on register button without entering name.", "Enter your username.", "Successful"),
                ("2.", "If user clicks on register button without entering phone number.", "Enter your phone number.", "Successful"),
                ("3.", "If user clicks on register button without entering email id.", "Enter your email id.", "Successful"),
                ("4.", "If user clicks on register button without entering password.", "Enter your password.", "Successful"),
                ("5.", "If user clicks on register button without entering confirm password.", "Enter your confirm password.", "Successful"),
                ("6.", "If password and confirm password mismatched.", "Password do not matched.", "Successful"),
                ("7.", "If the name field is filled in digits and clicks on register button.", "Name must contains only letters.", "Successful"),
                ("8.", "If the user enters phone number field less than or greater than 10 digits length.", "Phone number must be exactly 10 digits.", "Successful"),
                ("9.", "If the user enter invalid format of email id.", "Invalid email format.", "Successful"),
                ("10.", "If the user wants to register again with same phone number.", "This phone number is already registered.", "Successful"),
                ("11.", "If the user enter invalid format of password.", "Invalid password format.", "Successful"),
                ("12.", "If the valid register details are entered.", "System displays Login page.", "Successful"),
            ]
        },
        {
            "num": "5.6.2",
            "title": "User Login Form",
            "cases": [
                ("1.", "If user clicks login button without entering email/phone.", "Please enter email or phone number.", "Successful"),
                ("2.", "If user clicks login button without entering password.", "Please enter password.", "Successful"),
                ("3.", "If user enters unregistered email/phone.", "User not found with provided credentials.", "Successful"),
                ("4.", "If user enters incorrect password.", "Invalid credentials.", "Successful"),
                ("5.", "If user enters valid citizen credentials.", "System redirects to Citizen Dashboard.", "Successful"),
                ("6.", "If user enters valid official/contractor credentials.", "System redirects to respective Role Portal.", "Successful"),
                ("7.", "If user enters valid administrator credentials.", "System redirects to Admin Management Console.", "Successful"),
                ("8.", "If user clicks on 'Logout' option.", "User session cleared and redirects to Login.", "Successful"),
            ]
        },
        {
            "num": "5.6.3",
            "title": "Citizen Complaint Lodging (Live Camera & GPS)",
            "cases": [
                ("1.", "If user clicks submit complaint without selecting a park.", "Please select a target park.", "Successful"),
                ("2.", "If user submits complaint without selecting issue category.", "Please select an issue category.", "Successful"),
                ("3.", "If user submits complaint without entering issue description.", "Please provide issue description.", "Successful"),
                ("4.", "If citizen attempts to upload image from local gallery.", "Gallery upload disabled. Live camera capture mandatory.", "Successful"),
                ("5.", "If live camera captures geo-tagged image without device GPS enabled.", "Please enable device GPS / Location permissions.", "Successful"),
                ("6.", "If citizen completes live photo capture and submits valid complaint.", "Complaint registered successfully with unique tracking ID.", "Successful"),
                ("7.", "If citizen checks real-time status using complaint tracking ID.", "System displays current complaint status and workflow progress.", "Successful"),
            ]
        },
        {
            "num": "5.6.4",
            "title": "Admin Dual-Assignment & Complaint Management",
            "cases": [
                ("1.", "If admin views new complaints queue without assigning roles.", "Status remains 'Pending Assignment'.", "Successful"),
                ("2.", "If admin attempts to assign contractor without selecting an official.", "Dual-assignment requires both official and contractor.", "Successful"),
                ("3.", "If admin assigns contractor marked as 'On Leave'.", "Selected contractor is currently unavailable / on leave.", "Successful"),
                ("4.", "If admin assigns official marked as 'On Leave'.", "Selected official is currently unavailable / on leave.", "Successful"),
                ("5.", "If admin selects active official and active contractor with SLA deadline.", "Complaint status updated to 'Assigned' and notifications sent.", "Successful"),
                ("6.", "If admin updates/reassigns complaint after official rework request.", "Complaint updated and re-assigned to respective contractor.", "Successful"),
            ]
        },
        {
            "num": "5.6.5",
            "title": "Contractor Work Order & Proof Submission",
            "cases": [
                ("1.", "If contractor accesses work orders assigned to other contractors.", "Access denied. Only assigned orders visible.", "Successful"),
                ("2.", "If contractor submits completion report without uploading resolution photos.", "Mandatory to upload work completion proof images.", "Successful"),
                ("3.", "If contractor attempts to upload gallery image as work completion proof.", "Live camera capture with GPS timestamp required.", "Successful"),
                ("4.", "If contractor submits live work proof photos and remarks.", "Status changes to 'In Inspection' and notifies assigned official.", "Successful"),
                ("5.", "If contractor submits material requisition without required item list.", "Please specify required materials and quantities.", "Successful"),
                ("6.", "If contractor submits valid material requisition request.", "Material request sent to Admin for approval.", "Successful"),
            ]
        },
        {
            "num": "5.6.6",
            "title": "Official Physical Inspection & Digital Sign-off",
            "cases": [
                ("1.", "If official performs sign-off outside verified park geofence coordinates.", "Verification failed: Inspector must be within park geofence.", "Successful"),
                ("2.", "If official performs sign-off without taking live on-site inspection photo.", "Live on-site verification photo is required.", "Successful"),
                ("3.", "If official rejects work quality and requests rework.", "Status changed to 'Rework Required' and re-assigned to contractor.", "Successful"),
                ("4.", "If official approves work with remarks and digital sign-off.", "Complaint status updated to 'Resolved / Closed' and citizen notified.", "Successful"),
                ("5.", "If official submits leave application with valid dates.", "Leave request submitted to Admin for approval.", "Successful"),
            ]
        },
        {
            "num": "5.6.7",
            "title": "Citizen Park Facility Booking & Feedback",
            "cases": [
                ("1.", "If citizen books stall/event on already booked date slot.", "Selected slot is already reserved. Choose another date.", "Successful"),
                ("2.", "If citizen books stall without selecting category or purpose.", "Please select booking category and duration.", "Successful"),
                ("3.", "If citizen enters valid booking details and proceeds to payment.", "Booking confirmed and digital receipt generated.", "Successful"),
                ("4.", "If citizen submits feedback/rating for closed complaint.", "Feedback recorded and contractor performance rating updated.", "Successful"),
            ]
        },
        {
            "num": "5.6.8",
            "title": "Admin Master Data Management & Analytics",
            "cases": [
                ("1.", "If admin adds new park without mandatory geo-coordinates (lat/lng).", "Park latitude and longitude boundaries are required.", "Successful"),
                ("2.", "If admin registers new official/contractor account with missing fields.", "All mandatory personnel details must be filled.", "Successful"),
                ("3.", "If admin reviews and approves/rejects staff leave request.", "Leave status updated and user availability set accordingly.", "Successful"),
                ("4.", "If admin generates PDF/Excel audit and SLA compliance report.", "Audit report generated with compliant visual charts and metrics.", "Successful"),
            ]
        }
    ]

    for mod in modules:
        # Heading 3 format: e.g., 5.6.1 User Registration Form
        h_p = doc.add_paragraph()
        h_p.paragraph_format.space_before = Pt(14)
        h_p.paragraph_format.space_after = Pt(6)
        h_p.paragraph_format.keep_with_next = True
        
        r_num = h_p.add_run(f"{mod['num']}\t")
        r_num.font.name = 'Times New Roman'
        r_num.font.size = Pt(12)
        r_num.font.bold = True
        r_num.font.color.rgb = RGBColor(0, 0, 0)
        
        r_title = h_p.add_run(f"{mod['title']}")
        r_title.font.name = 'Times New Roman'
        r_title.font.size = Pt(12)
        r_title.font.bold = True
        r_title.font.color.rgb = RGBColor(0, 0, 0)
        
        # Table
        table = doc.add_table(rows=1 + len(mod['cases']), cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_table_borders(table)
        
        col_widths = [Inches(0.7), Inches(2.7), Inches(2.2), Inches(1.0)]
        
        headers = ["Sl. No.", "Test Condition", "Expected Result", "Result"]
        hdr_row = table.rows[0]
        hdr_row._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        
        for i, title in enumerate(headers):
            cell = hdr_row.cells[i]
            cell.width = col_widths[i]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i in [0, 3] else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            
            r = p.add_run(title)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0, 0, 0)
            
        for row_idx, case_data in enumerate(mod['cases']):
            row = table.rows[row_idx + 1]
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            
            for col_idx in range(4):
                cell = row.cells[col_idx]
                cell.width = col_widths[col_idx]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx in [0, 3] else WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                
                r = p.add_run(case_data[col_idx])
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor(0, 0, 0)
                
        # Spacing after table
        sp_p = doc.add_paragraph()
        sp_p.paragraph_format.space_before = Pt(4)
        sp_p.paragraph_format.space_after = Pt(6)

    doc.save(filename)
    print(f"File successfully created at: {filename}")

if __name__ == "__main__":
    create_test_cases_docx(r"d:\park_monitoring_system08\park-management\System_Test_Cases_All.docx")
