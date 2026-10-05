import os
import docx
from docx.shared import Pt, RGBColor

DOC_PATH_3 = r'd:\park_monitoring_system08\park-management\exact_pdf_docx_chapters\03_Chapter_3_System_Analysis.docx'

features_text = [
    ("1. Public User / Citizen", "Heading 2"),
    ("Park Discovery & Amenity Search: Browse municipal parks by Zone/Ward, view operational timings, facility lists (benches, lights, walking tracks, children play area, restrooms), and green ratings.", "Normal"),
    ("Camera & GPS Issue Reporting: Report maintenance complaints (broken equipment, cleanliness, lighting failure) with live photo capture and device GPS coordinate logging.", "Normal"),
    ("Complaint Tracking: Live tracking of reported issues (Submitted → Assigned → In-Progress → Resolved).", "Normal"),
    ("Commercial Stall Booking: Reserve temporary commercial stall slots in designated park zones with real-time slot selection and digital payment receipts.", "Normal"),
    ("Eco & Community Events: View and register for upcoming park events (tree plantation drives, nature walks, yoga sessions).", "Normal"),
    ("Feedback & Rating: Submit ratings and qualitative reviews on park maintenance and cleanliness.", "Normal"),
    ("2. Municipal Administrator (Admin)", "Heading 2"),
    ("Master Data Management: Manage municipal zones, divisions, wards, and park records.", "Normal"),
    ("Inspection Task Assignment: Select a specific park and assign an on-demand inspection task to a Govt Official.", "Normal"),
    ("Contractor Work Assignment: Review reported park complaints and assign repair jobs to specialized Contractors (Electrician, Plumber, Mason, Carpenter, General Maintenance).", "Normal"),
    ("Stall & Event Management: Approve/manage commercial stall bookings and schedule eco-community events.", "Normal"),
    ("Analytics & Reports: Monitor resolution SLAs, contractor turnaround times, and park status across all zones.", "Normal"),
    ("3. Government Official (Field Inspector)", "Heading 2"),
    ("Assigned Park Inspections: Receive inspection assignments from the Admin, visit the park, check physical amenities, and submit the inspection report.", "Normal"),
    ("Work Verification & Final Sign-Off: Inspect on-site repair work completed by contractors, verify photo proof, and give the official approval/sign-off to close the complaint.", "Normal"),
    ("Work Schedule & Notifications: Manage inspection schedules and view assignment alerts.", "Normal"),
    ("4. Contractor (Maintenance Provider)", "Heading 2"),
    ("Task Dashboard: View repair jobs assigned by the Admin with priority, park location, and complaint photos.", "Normal"),
    ("Work Progress Updates: Update job status (Accepted → In-Progress → Completed).", "Normal"),
    ("Completion Proof Submission: Upload photos and description of completed repair work for Govt Official review and approval.", "Normal"),
    ("Material Requests: Submit requests for spare parts/materials required for maintenance.", "Normal")
]

def append_features(doc_path):
    doc = docx.Document(doc_path)
    
    p = doc.add_paragraph()
    p.style = 'Heading 1'
    r = p.add_run("Detailed Project Features & Modules")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(0, 51, 102)
    
    for text, style in features_text:
        p = doc.add_paragraph()
        if style == "Heading 2":
            p.style = 'Heading 2'
            r = p.add_run(text)
            r.bold = True
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(0, 51, 102)
        else:
            p.style = 'Normal'
            if ':' in text:
                title, desc = text.split(':', 1)
                r1 = p.add_run(title + ':')
                r1.bold = True
                p.add_run(desc)
            else:
                p.add_run(text)
                
    doc.save(doc_path)
    print(f"Features appended to {doc_path}")

try:
    append_features(DOC_PATH_3)
except Exception as e:
    print(e)
