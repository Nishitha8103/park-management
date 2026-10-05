import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

OUTPUT_DIR = r"d:\park_monitoring_system08\park-management\exact_pdf_docx_chapters"
COMBINED_FILE = os.path.join(OUTPUT_DIR, "Complete_Park_Monitoring_System_Report.docx")

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_full_report():
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1)
        
        header = section.header
        hp = header.paragraphs[0]
        hp.text = " Park Monitoring System"
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hp.style.font.name = 'Times New Roman'
        hp.style.font.size = Pt(11)
        hp.style.font.bold = True
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="double" w:sz="12" w:space="4" w:color="000000"/></w:pBdr>')
        hp._element.get_or_add_pPr().append(pBdr)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = "Department of MCA\t\tPIM, Udupi"
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fp.style.font.name = 'Times New Roman'
        fp.style.font.size = Pt(11)
        fp.style.font.bold = True
        pBdrF = parse_xml(f'<w:pBdr {nsdecls("w")}><w:top w:val="double" w:sz="12" w:space="4" w:color="000000"/></w:pBdr>')
        fp._element.get_or_add_pPr().append(pBdrF)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_after = Pt(6)

    def add_h1(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(12)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(15)
        run.font.bold = True

    def add_h2(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        run.font.bold = True

    def add_h3(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.font.bold = True

    def add_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(6)
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.bold = True
            rb.font.name = 'Times New Roman'
            rb.font.size = Pt(12)
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)

    def add_bullet_text(text, bold_title=None):
        p = doc.add_paragraph(style='List Bullet')
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(4)
        if bold_title:
            rb = p.add_run(bold_title + "\n" if not bold_title.endswith(":") else bold_title + " ")
            rb.font.bold = True
            rb.font.name = 'Times New Roman'
            rb.font.size = Pt(12)
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)

    def create_table(headers, rows):
        t = doc.add_table(rows=1, cols=len(headers))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(headers):
            t.rows[0].cells[i].text = h
            t.rows[0].cells[i].paragraphs[0].runs[0].font.bold = True
            set_cell_background(t.rows[0].cells[i], "E2E8F0")
        for r_data in rows:
            r = t.add_row().cells
            for idx, val in enumerate(r_data):
                r[idx].text = val
        doc.add_paragraph()

    # Title Page
    add_h1("MANGALORE UNIVERSITY")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Project Report on\n“PARK MONITORING SYSTEM”\n\nCarried out and submitted by\n")
    r.font.bold = True
    r.font.size = Pt(14)
    r2 = p.add_run("NISHITHA\nReg. No.: P05PP23S126010\nIV Semester MCA Student\nPoornaprajna Institute of Management, Udupi\n\n")
    r2.font.bold = True
    r2.font.size = Pt(13)
    r3 = p.add_run("Under the guidance of\nProf. Venugopala Rao A. S.\nHOD\nDept.of MCA, PIM, Udupi.\n\nIn the partial fulfilment of the requirement for the award of degree in Master of Computer Applications during the academic year 2024-25\n\nPoornaprajna Institute of Management,\nUdupi\n2024-2025\n")
    r3.font.size = Pt(12)

    # Certificate
    doc.add_page_break()
    add_h1("CERTIFICATE")
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("MANGALORE UNIVERSITY\nPoornaprajna Institute of Management, Udupi\nDepartment of MCA\n")
    r_inst.font.bold = True
    r_inst.font.size = Pt(13)
    add_p("This is to certify that the project entitled “Park Monitoring System” has been carried out by Nishitha (Reg.No.: P05PP23S126010), student of fourth semester MCA (Masters of Computer Applications) under the supervision of Prof. Venugopala Rao A.S., MCA Department, Poornaprajna Institute of Management, Udupi. The project is submitted in partial fulfilment of the requirement for the award of Master of Computer Applications by Mangalore University during the Academic year 2024-2025.\n\n\n")
    p_sig = doc.add_paragraph()
    p_sig.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_sig.add_run("Internal Guide\t\t\t\t\tHead of the Department\n\n\n\nInternal Examiner\t\t\t\t\tExternal Examiner\n\nSubmitted for the viva-voice examination held on: ..................")

    # Declaration
    doc.add_page_break()
    add_h1("DECLARATION")
    add_p("This project work entitled “PARK MONITORING SYSTEM” has been successfully carried out by me under the supervision and guidance of Prof. Venugopala Rao A.S., Head of the MCA Department, Poornaprajna Institute of Management, Udupi. This project is submitted in partial fulfilment for the award of Masters of Computer Application degree by Mangalore University during the academic year 2024–25.\n\nThis work or any part of this work has not been submitted to any other university or Institute/School for the award of any other Degree or Diploma.\n\n\nDate:\t\t\t\t\t\tName: NISHITHA\nPlace: Udupi\t\t\t\t\tReg. No.: P05PP23S126010")

    # Acknowledgement
    doc.add_page_break()
    add_h1("ACKNOWLEDGEMENT")
    add_p("I take this opportunity to express my sincere thanks to my guide Prof. Venugopala Rao A. S, Head of the MCA Department, Poornaprajna Institute of Management, for him all-round guidance, timely help at every stage of this project and for the valuable suggestions and unlimited support.")
    add_p("I would like to express my gratitude to Dr. P. S. Aithal, Director, Poornaprajna Institute of Management, for granting the necessary permissions and providing all kinds of infrastructure facilities in the department to carry out this project successfully.")
    add_p("All faculty members and non-teaching staffs of the MCA department.")
    add_p("I would like to express my heartfelt gratitude to my parents, friends and well-wishers who have always inspired and blessed me, including those whom I may have inadvertently failed to mention. Above all, with all my heart, I thank you God for empowering me with the dedication, focus and patience to carry out this project successfully.\n\n\nNISHITHA\nP05PP23S126010")

    # Abstract
    doc.add_page_break()
    add_h1("ABSTRACT")
    add_p("Urban green space and public park management is one of the critical challenges faced by rapidly growing cities like Bangalore. The increasing wear and tear of park amenities, damaged play equipment, broken benches, lighting failures, and neglected green zones have become a major concern for the Bruhat Bengaluru Mahanagara Palike (BBMP). These neglected park issues not only affect the city’s cleanliness and natural beauty but also pose safety risks to visiting citizens and children. To address these issues, there is a pressing need for a digital platform that enables timely reporting, monitoring, and resolution of such problems through citizen participation and administrative coordination. This project presents the development of a Progressive Web Application (PWA) aimed at supporting BBMP’s park monitoring and maintenance initiative. The system provides a real-time, transparent, and user-friendly platform for reporting, tracking, and resolving park maintenance complaints across Bengaluru. It bridges the communication gap between Citizens, Field Officials, Maintenance Contractors, and Administrators by ensuring smooth information flow and coordinated action.")
    add_p("The system comprises four major modules: Citizens, Admin, Government Officials, and Maintenance Contractors. The Citizens module enables citizens to register, log in, update profiles, view park directories, book stall slots with payments, register for events, and raise complaints by capturing live photos through their camera, location details, and descriptions of damaged park facilities. Citizens can also check the live status of their complaints. Additionally, the system sends email notifications to citizens upon successful complaint submission, status updates, and resolution, ensuring transparency and keeping citizens informed throughout the process.")
    add_p("The Admin module acts as the central control panel, managing all activities such as complaint assignments, user and official management, contractor monitoring, park asset records, stall bookings, and report generation. It includes sub modules for managing users, officials, contractors, districts, corporations, zones, and wards, allowing the administrator to maintain proper geographic and operational control. When a complaint is assigned or resolved, automated email alerts are triggered to both the assigned field staff and the concerned citizen, improving communication efficiency.")
    add_p("The Field Officials and Contractors module empowers on-ground staff to log in, view assigned complaints, verify reported defects, log daily work progress, upload proof of resolution, request materials, and maintain their work history for performance tracking. By integrating these modules with real-time email notification support, the application ensures seamless communication, efficient complaint resolution, and improved park management workflows.")
    add_p("Keywords: Park Monitoring, PWA, Proactive, Transparency, Complaints", "")

    # Table of Contents
    doc.add_page_break()
    add_h1("Table of Contents")
    create_table(["Chapter No.", "Chapter Name", "Page No."], [
        ("1.", "Introduction\n1.1 Introduction\n1.2 Overview of the project\n1.3 Problem Statement\n1.4 Motivation\n1.5 Significance of the Study\n1.6 Objectives\n1.7 Scope of the Project\n1.8 Features", "1-11"),
        ("2.", "Literature Review\n2.1 Introduction\n2.2 Result analysis\n2.3 Identified gaps in the Literature\n2.4 Existing System\n2.5 Proposed System", "12-16"),
        ("3.", "System Analysis\n3.1 Introduction\n3.2 Overall Descriptions\n3.3 Specific Requirements\n3.4 Functional Requirements\n3.5 Performance Requirements\n3.6 Design Constraints\n3.7 Other Requirements\n3.8 Safety Requirements\n3.9 Security Requirements", "17-23"),
        ("4.", "Design And Methodology\n4.1 System Design\n4.2 Detailed Design\n4.3 Database Design", "24-36"),
        ("5.", "Implementation Details\n5.1 Introduction\n5.2 Hardware And Software tools used\n5.3 Source Code", "37-46"),
        ("6.", "Result and Evaluation\n6.1 Introduction\n6.2 Test Scenario\n6.3 Test Cases", "47-51"),
        ("7.", "Conclusion and Future Work\n7.1 Conclusion\n7.2 Future Work", "52-53"),
        ("", "Appendices", "54-64"),
        ("", "References", "65")
    ])

    # List of Figures
    doc.add_page_break()
    add_h1("List of Figures")
    create_table(["Figure No.", "Figure Name", "Figure Page No."], [
        ("4.1.2.1.1", "Use Case Diagram for Citizens", "26"),
        ("4.1.2.1.2", "Use Case Diagram for Admin", "26"),
        ("4.1.2.1.3", "Use Case Diagram for Field Officials", "27"),
        ("4.1.2.2", "Context Flow Diagram", "27"),
        ("4.2.1.1", "Data Flow Diagram for Citizens", "29"),
        ("4.2.1.2", "Structure Chart for Admin", "31"),
        ("4.2.1.3", "UML Diagram for Field Officials", "33"),
        ("Fig.1 to Fig.30", "Fig:1 Index Page, Fig:2 Register\nFig:3 Login, Fig:4 Forgot Password\nFig:5 OTP Page, Fig:6 Reset ,Fig:7 Citizen Dashboard,Fig:8 Profile, Fig:9 Edit Profile,Fig:10 Add Complaints,\nFig:11 Check Status,Fig:12 Complaint List, Fig:13 Closed Complaint List,\nFig:14 Admin Login,Fig:15 Admin Dashboard,Fig:16Report,Fig:17Compalint\nFig:18 Add Zone, Fig:19 Division,Fig:20 Wards,Fig:21 Add Official,Fig:22 Official Dashboard, Fig:23 Assign Compalints,Fig:24 History, Fig:25 Update\nFig:26 Complaint Model,Fig:27 Officials Model,Fig:28 Ward Table,Fig: 29 Division Table,Fig: 30 Complaint assignment.", "54-64")
    ])

    # List of Tables
    doc.add_page_break()
    add_h1("List of Tables")
    create_table(["Table No.", "Table Name", "Page No."], [
        ("4.3.1.1", "Users", "34"),
        ("4.3.1.2", "Admin", "34"),
        ("4.3.1.3", "Complaints", "35"),
        ("4.3.1.4", "Complaint_Officials", "35"),
        ("4.3.1.5", "Zones", "35"),
        ("4.3.1.6", "Divisions", "36"),
        ("4.3.1.7", "Wards", "36"),
        ("4.3.1.8", "Officials", "36"),
        ("6.3.1", "Registration Form", "48"),
        ("6.3.2", "Login", "49"),
        ("6.3.3", "Forgot Password", "49"),
        ("6.3.4", "Raise Complaints", "49"),
        ("6.3.5", "Check status", "49"),
        ("6.3.6", "Report", "50"),
        ("6.3.7", "Add zones", "50"),
        ("6.3.8", "Add Divisions", "50"),
        ("6.3.9", "Add wards", "51"),
        ("6.3.10", "Add Officials", "51")
    ])

    doc.save(COMBINED_FILE)
    print("Generated Consolidated Bound Report:", COMBINED_FILE)

if __name__ == "__main__":
    create_full_report()
