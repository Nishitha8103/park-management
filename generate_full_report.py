import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

OUTPUT_DIR = r"d:\park_monitoring_system08\park-management\documentation_chapters"
COMBINED_FILE = os.path.join(OUTPUT_DIR, "Complete_Park_Monitoring_Project_Report.docx")

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_full_report():
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1)
        
        header = section.header
        hp = header.paragraphs[0]
        hp.text = "Park Monitoring & Management System"
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hp.style.font.name = 'Times New Roman'
        hp.style.font.size = Pt(9)
        hp.style.font.color.rgb = RGBColor(120, 120, 120)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.text = "Department of MCA, PIM, Udupi"
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp.style.font.name = 'Times New Roman'
        fp.style.font.size = Pt(9)
        fp.style.font.color.rgb = RGBColor(120, 120, 120)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(30, 30, 30)
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_after = Pt(6)

    def add_h1(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(12)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 81, 50)

    def add_h2(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(21, 128, 61)

    def add_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(6)
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.bold = True
            rb.font.name = 'Times New Roman'
            rb.font.size = Pt(12)
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)

    def add_bullet(title, desc):
        p = doc.add_paragraph(style='List Bullet')
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(4)
        rb = p.add_run(title + ": ")
        rb.font.bold = True
        rb.font.name = 'Times New Roman'
        rb.font.size = Pt(12)
        r = p.add_run(desc)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)

    # ------------------ PRELIMINARY ------------------
    add_h1("MANGALORE UNIVERSITY")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Project Report on\n“PARK MONITORING & MANAGEMENT SYSTEM”\n\nCarried out and submitted by\n")
    r.font.bold = True
    r.font.size = Pt(14)
    r2 = p.add_run("NISHITHA\nReg. No.: P05PP23S126010\nIV Semester MCA Student\nPoornaprajna Institute of Management, Udupi\n\n")
    r2.font.bold = True
    r2.font.size = Pt(13)
    r3 = p.add_run("Under the guidance of\nProf. Venugopala Rao A. S.\nHOD, Dept. of MCA, PIM, Udupi.\n\nIn the partial fulfilment of the requirement for the award of degree in Master of Computer Applications during the academic year 2024-2025\n\nPoornaprajna Institute of Management, Udupi\n2024-2025\n")
    r3.font.size = Pt(12)

    doc.add_page_break()
    add_h1("CERTIFICATE")
    add_p("This is to certify that the project entitled “Park Monitoring & Management System” has been carried out by Nishitha (Reg.No.: P05PP23S126010), student of fourth semester MCA (Masters of Computer Applications) under the supervision of Prof. Venugopala Rao A.S., MCA Department, Poornaprajna Institute of Management, Udupi. The project is submitted in partial fulfilment of the requirement for the award of Master of Computer Applications by Mangalore University during the Academic year 2024-2025.\n\n\n")
    p_sig = doc.add_paragraph()
    p_sig.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_sig.add_run("Internal Guide\t\t\t\t\tHead of the Department\n\n\n\nInternal Examiner\t\t\t\t\tExternal Examiner\n\nSubmitted for the viva-voce examination held on: ..................")

    doc.add_page_break()
    add_h1("DECLARATION")
    add_p("This project work entitled “PARK MONITORING & MANAGEMENT SYSTEM” has been successfully carried out by me under the supervision and guidance of Prof. Venugopala Rao A.S., Head of the MCA Department, Poornaprajna Institute of Management, Udupi. This project is submitted in partial fulfilment for the award of Masters of Computer Application degree by Mangalore University during the academic year 2024–25. This work or any part of this work has not been submitted to any other university or Institute/School for the award of any other Degree or Diploma.\n\n\nDate: ...................\nPlace: Udupi\t\t\t\t\tName: NISHITHA\n\t\t\t\t\t\tReg. No.: P05PP23S126010")

    doc.add_page_break()
    add_h1("ACKNOWLEDGEMENT")
    add_p("I take this opportunity to express my sincere thanks to my guide Prof. Venugopala Rao A. S., Head of the MCA Department, Poornaprajna Institute of Management, for his all-round guidance, timely help at every stage of this project, and for valuable suggestions and unlimited support.\n\nI would like to express my gratitude to Dr. P. S. Aithal, Director, Poornaprajna Institute of Management, for granting the necessary permissions and providing all kinds of infrastructure facilities in the department to carry out this project successfully.\n\nI also express my sincere gratitude to all faculty members and non-teaching staff of the MCA department for their constant encouragement.\n\nI would like to express my heartfelt gratitude to my parents, friends, and well-wishers who have always inspired and supported me. Above all, with all my heart, I thank God for empowering me with the dedication, focus, and patience to carry out this project successfully.\n\n\nNISHITHA\nReg. No.: P05PP23S126010")

    doc.add_page_break()
    add_h1("ABSTRACT")
    add_p("Urban recreational and ecological spaces such as public parks are fundamental to sustainable urban living, biodiversity, air quality, and public health in rapidly growing metropolitan cities like Bengaluru. However, maintaining large municipal park infrastructures across hundreds of city wards is a formidable civic challenge for municipal corporations like Bruhat Bengaluru Mahanagara Palike (BBMP). Traditional manual inspections and paper-based complaint registers frequently suffer from delayed grievance reporting, uncoordinated maintenance workflows, lack of transparent tracking, and underutilized public engagement.")
    add_p("To overcome these challenges, this project presents the 'Park Monitoring & Management System', an advanced, end-to-end Progressive Web Application (PWA) built with the MERN Stack (MongoDB, Express.js, React, Node.js). The platform provides unified, multi-role digital coordination across four specialized stakeholder modules: Public Citizens, Government Civic Officials, Maintenance Contractors, and Municipal Super Administrators.")
    add_p("Citizens can explore geo-mapped parks, verify amenities, register complaints with mandatory live camera capture and real-time GPS coordinates, book commercial/community stall spaces with instant QR digital payments, register for park eco-events, and submit structured feedback. Municipal Administrators benefit from real-time analytics dashboards, automated SLA compliance tracking, multi-district geographic hierarchy (District → Corporation → Zone → Ward), and contractor task dispatching. Government Officials inspect work quality, review before/after photo evidence, manage emergency alerts, and verify maintenance. Maintenance Contractors receive automated task tickets, submit daily work progress logs, request raw materials, and file work completion reports.")
    add_p("By bridging the communication gap between citizens and civic authorities through automated notifications, QR code park verification, and geospatial telemetry, the system fosters transparent urban governance and preserves vibrant urban green cover.", "Keywords: ")
    add_p("Park Monitoring, Progressive Web Application (PWA), MERN Stack, Geolocation Tracking, SLA Compliance, Municipal Governance, Stall Booking, Complaint Redressal.")

    # Table of Contents
    doc.add_page_break()
    add_h1("TABLE OF CONTENTS")
    toc_data = [
        ("1.", "Introduction", "1"),
        ("", "1.1 Introduction", "1"),
        ("", "1.2 Overview of the Project", "2"),
        ("", "1.3 Problem Statement", "3"),
        ("", "1.4 Motivation", "4"),
        ("", "1.5 Significance of the Study", "5"),
        ("", "1.6 Objectives of the Project", "7"),
        ("", "1.7 Scope of the Project", "8"),
        ("", "1.8 Features of the Application", "10"),
        ("2.", "Literature Review", "12"),
        ("", "2.1 Introduction", "12"),
        ("", "2.2 Result Analysis", "12"),
        ("", "2.3 Identified Gaps in Literature", "14"),
        ("", "2.4 Existing System", "15"),
        ("", "2.5 Proposed System", "16"),
        ("3.", "System Analysis", "17"),
        ("", "3.1 Introduction", "17"),
        ("", "3.2 Overall Description", "18"),
        ("", "3.3 Specific Requirements", "19"),
        ("", "3.4 Functional Requirements", "20"),
        ("", "3.5 Non-Functional & Performance Requirements", "22"),
        ("4.", "Design and Methodology", "24"),
        ("", "4.1 System Design Principles", "24"),
        ("", "4.2 High-Level Architectural Flow", "26"),
        ("", "4.3 Database Schema & Collection Design", "28"),
        ("", "4.4 UML Design Specifications", "31"),
        ("5.", "Implementation Details", "37"),
        ("", "5.1 Introduction", "37"),
        ("", "5.2 Technological Stack & Tools", "37"),
        ("", "5.3 Core Source Code Implementations", "39"),
        ("6.", "Result and Evaluation", "47"),
        ("", "6.1 Introduction", "47"),
        ("", "6.2 Test Scenarios & Methodology", "47"),
        ("", "6.3 Comprehensive Test Cases & Results", "48"),
        ("", "6.4 Performance & Usability Evaluation", "51"),
        ("7.", "Conclusion and Future Work", "52"),
        ("", "7.1 Conclusion", "52"),
        ("", "7.2 Future Work", "53"),
        ("", "References", "54"),
        ("", "Appendices (Screenshots & Navigations)", "56"),
    ]
    t_table = doc.add_table(rows=1, cols=3)
    t_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_table.rows[0].cells[0].text = "Chapter No."
    t_table.rows[0].cells[1].text = "Chapter / Section Name"
    t_table.rows[0].cells[2].text = "Page No."
    for cell in t_table.rows[0].cells:
        cell.paragraphs[0].runs[0].font.bold = True
        set_cell_background(cell, "D1E7DD")

    for row in toc_data:
        r_cells = t_table.add_row().cells
        r_cells[0].text = row[0]
        r_cells[1].text = row[1]
        r_cells[2].text = row[2]

    # Chapter 1
    doc.add_page_break()
    add_h1("CHAPTER 1\nINTRODUCTION")
    add_h2("1.1 Introduction")
    add_p("Urbanization has brought unprecedented growth, industrial modernization, and economic vitality to metropolitan cities. However, this swift expansion has exerted severe pressure on natural ecosystems and public open spaces. In modern urban ecosystems, public parks, botanical reserves, and neighborhood green lungs play an irreplaceable role in preserving biodiversity, mitigating the urban heat island effect, purifying ambient air, and supporting the physical and psychological well-being of city residents.")
    add_p("Bengaluru, renowned as the 'Garden City of India', hosts an extensive network of thousands of public parks managed under the civic umbrella of the Bruhat Bengaluru Mahanagara Palike (BBMP). Maintaining this massive green infrastructure—encompassing walking tracks, children's play equipment, open-air gymnasiums, solar lighting, sanitation facilities, lakes, and botanical flora—presents an intricate administrative and operational challenge.")
    add_p("Traditional park administration in municipal bodies has historically relied on physical inspections, manual paperwork, decentralized record-keeping, and informal complaint reporting. These legacy workflows frequently encounter severe bottlenecks: complaints regarding broken play equipment, damaged benches, overflowing garbage bins, broken lighting, or dried horticulture often go unattended for weeks due to miscommunication between field supervisors, contractors, and administrative headquarters. Moreover, citizens lack a transparent window into park amenities, ongoing maintenance schedules, and public event reservations.")
    add_p("In this context, the 'Park Monitoring & Management System' is developed as a modern, full-stack Progressive Web Application (PWA). It unifies citizen engagement, contractor maintenance execution, official inspections, and administrative governance into a single, high-performance cloud platform.")

    add_h2("1.2 Overview of the Project")
    add_p("The Park Monitoring & Management System is an enterprise-grade civic management solution designed to digitize, streamline, and automate the governance of urban parks. By harnessing the MERN stack (MongoDB, Express.js, React.js, Node.js) paired with Progressive Web App capabilities, the system provides an ultra-responsive, app-like experience on smartphones, tablets, and desktop workstations without requiring complex app-store installations.")
    add_p("The platform is architected into four interconnected operational modules:")
    add_bullet("1. Citizen & Public Portal", "Empowers residents to locate parks on interactive Leaflet maps, filter by specific amenities (yoga lawn, open gym, wheelchair ramps), submit geo-tagged complaints with live camera proof, book authorized commercial/food stall slots with integrated online payments, reserve tickets for park cultural/environmental events, and trigger one-touch Emergency SOS alerts.")
    add_bullet("2. Government Officials Module", "Equips municipal ward engineers, horticultural officers, and civic inspectors with mobile-optimized dashboards to inspect assigned parks, track daily schedules, review contractor progress logs, verify before/after cleanup imagery, and validate work completion with geo-auditing.")
    add_bullet("3. Maintenance Contractors Module", "Enables specialized maintenance agencies to view assigned work tickets, submit real-time milestone progress updates, upload timestamped completion photographs, raise formal material requests, and manage workforce allocations.")
    add_bullet("4. Central Super Admin Control Panel", "Provides municipal administrators with end-to-end command over the spatial hierarchy (Districts, Corporations, Zones, Wards), dynamic park asset registrations, contractor KYC reviews, automated SLA tracking, stall fee revenues, and automated broadcast announcements.")

    add_h2("1.3 Problem Statement")
    add_p("Urban municipal authorities face numerous persistent hurdles in managing civic green spaces efficiently:")
    add_bullet("Lack of Centralized Digital Inventory", "Parks, horticulture assets, lighting fixtures, and gym equipment are recorded in disparate ward logs, preventing holistic inventory visibility.")
    add_bullet("Opaque Complaint Redressal", "Citizens reporting broken amenities or safety hazards through physical registers or social media have no visibility into ticket progress, resolution timelines, or assigned contractors.")
    add_bullet("Absence of Verification Proof", "Contractors often report maintenance completion without verifiable photographic proof or geotagged inspection stamps, leading to substandard upkeep.")
    add_bullet("Unregulated Commercial Stalls & Events", "Informal park vending leads to encroachment and revenue leakage for municipal corporations due to the absence of a transparent digital stall booking and payment infrastructure.")
    add_bullet("Emergency Response Gaps", "Visitors facing medical emergencies or security threats in sprawling parks lack an instant mechanism to transmit precise coordinate-based distress signals to ward rangers.")

    add_h2("1.4 Motivation")
    add_p("The rapid growth of Bengaluru demands smart, sustainable, and data-driven governance. Public parks are essential community assets that directly influence the health and happiness of citizens. The motivation behind this project stems from the desire to leverage modern cloud and web technologies to make urban park management proactive, transparent, and participatory.")
    add_p("Empowering citizens with an instant, mobile-friendly reporting mechanism and furnishing municipal officers with automated task verification tools transforms park maintenance from a reactive chore into an efficient, continuous service.")

    add_h2("1.5 Significance of the Study")
    add_p("This project represents a crucial contribution to the Smart Cities Mission and Digital India e-governance paradigms:")
    add_bullet("Societal Significance", "Fosters civic ownership, enhances community safety, promotes green lifestyle activities, and builds democratic transparency between citizens and civic bodies.")
    add_bullet("Administrative Significance", "Replaces error-prone paperwork with automated SLA timers, role-based access control, automated email notifications, and auditable contractor performance scores.")
    add_bullet("Environmental & Ecological Significance", "Ensures swift resolution of horticultural blight, broken irrigation, and littering, safeguarding urban flora, fauna, and water bodies.")

    add_h2("1.6 Objectives of the Project")
    add_p("The primary objectives of the Park Monitoring & Management System are:")
    add_bullet("1. Develop a Universal Progressive Web App (PWA)", "Deliver a high-performance, mobile-responsive web application accessible across iOS, Android, and Desktop platforms.")
    add_bullet("2. Implement Geo-Tagged Live Camera Complaint Filing", "Enforce live camera capture and GPS geolocation to prevent fraudulent grievance filings and ensure accurate ground locating.")
    add_bullet("3. Build an Automated SLA & Contractor Tracking System", "Track complaint turnaround times against statutory service-level agreements (On Time, Due Soon, Overdue) and automate contractor dispatch.")
    add_bullet("4. Digitize Park Stall Booking & Event Registrations", "Provide citizens and local vendors with dynamic slot availability, transparent fee structures, and automated payment receipts.")
    add_bullet("5. Provide Role-Based Interactive Dashboards", "Deliver tailored interfaces with custom palettes (Dark Professional for Admin, Olive Sage for Officials, Vibrant Green for Citizens) complete with real-time KPI analytics and interactive Leaflet maps.")

    add_h2("1.7 Scope of the Project")
    add_p("The scope spans complete administrative, functional, and technological domains for municipal civic management:")
    add_bullet("Functional Scope", "Includes User Registration/Login, Interactive Park Directory, Facility Matrix, Live Camera Complaint Redressal, Contractor Milestone Updates, Official Physical Verifications, Stall Slot Matrix, Event Registration Engine, Automated Email Broadcasts, and SweetAlert2 UX.")
    add_bullet("Geographical Scope", "Structured to scale across Districts, Corporations (e.g., BBMP), Zones (East, West, South, Yelahanka, Mahadevapura, etc.), and individual Wards.")
    add_bullet("Technical Scope", "MERN Stack (MongoDB, Express, React 19, Node.js), Vite bundler, Leaflet.js mapping, Recharts visual analytics, JWT Authentication, and Web Push / Email notification dispatchers.")

    add_h2("1.8 Key Features of the Application")
    add_bullet("Progressive Web App (PWA) Engine", "Full offline asset caching, home screen installability, and lightning-fast loading via service workers.")
    add_bullet("Live Camera Mandatory Capture", "Direct hardware camera access for grievance photos, eliminating stale gallery uploads.")
    add_bullet("Hierarchical Spatial Master Setup", "Cascading administrative filters from District down to individual Parks.")
    add_bullet("Interactive Leaflet Geospatial Maps", "Real-time health visualization of parks with color-coded status pins.")
    add_bullet("Automated SLA Tracking & Notifications", "Proactive ticket lifecycle tracking with instant transactional email dispatches.")

    # Save
    doc.save(COMBINED_FILE)
    print("Generated Consolidated Report:", COMBINED_FILE)

if __name__ == "__main__":
    create_full_report()
